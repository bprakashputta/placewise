"""Block 1: rendering PDF / DOCX / image PDF / corrupt DOCX. Writes only to pytest tmp dirs."""
import io
import re
import zipfile
import zlib

import pytest
from reportlab.pdfgen import canvas as rl_canvas

from placewise.seed import render as R
from placewise.seed.noise import INJECTION_TEXT

from seed_helpers import make_facts, make_resume


@pytest.fixture
def fx():
    c = make_facts()
    r = make_resume(c)
    return c, r.model_dump(), r.to_text()


def _pdf_text_streams(data: bytes) -> bytes:
    """Concatenate all content streams, inflating FlateDecode ones."""
    import base64
    out = []
    for m in re.finditer(rb"/Filter\s*\[?([^\]>]*)\]?.*?stream\r?\n(.*?)endstream", data, re.S):
        filters, raw = m.group(1), m.group(2).strip()
        try:
            if b"ASCII85Decode" in filters:
                raw = base64.a85decode(raw, adobe=raw.startswith(b"<~") or raw.endswith(b"~>"))
            if b"FlateDecode" in filters:
                raw = zlib.decompress(raw)
        except Exception:
            pass
        out.append(raw)
    return b"\n".join(out)


@pytest.mark.parametrize("template", list(R.THEMES))
def test_render_pdf_every_template(fx, template):
    c, r, _ = fx
    data = R.render_pdf(r, c, template)
    assert data.startswith(b"%PDF") and data.rstrip().endswith(b"%%EOF")


@pytest.mark.parametrize("template", list(R.THEMES))
def test_render_docx_every_template_is_valid_zip(fx, template):
    c, r, _ = fx
    data = R.render_docx(r, c, template)
    assert zipfile.is_zipfile(io.BytesIO(data))
    from docx import Document
    doc = Document(io.BytesIO(data))
    body = "\n".join(p.text for p in doc.paragraphs) + "".join(
        cell.text for t in doc.tables for row in t.rows for cell in row.cells)
    assert "Testy Fixture" in body and "P.E., Texas #123456" in body


def test_render_image_pdf(fx):
    c, _, text = fx
    data = R.render_image_pdf(text, c)
    assert data.startswith(b"%PDF")
    assert b"/Image" in data
    assert b"Testy Fixture" not in _pdf_text_streams(data)  # no text layer


def test_injection_text_drawn_only_when_flagged(fx, monkeypatch):
    c, r, _ = fx
    drawn = []
    orig = rl_canvas.Canvas.drawString

    def spy(self, x, y, text, *a, **kw):
        drawn.append(text)
        return orig(self, x, y, text, *a, **kw)

    monkeypatch.setattr(rl_canvas.Canvas, "drawString", spy)
    clean = R.render_pdf(r, c, "classic")
    assert INJECTION_TEXT not in drawn
    c.injection = True
    data = R.render_pdf(r, c, "classic")
    assert INJECTION_TEXT in drawn
    assert data.startswith(b"%PDF") and len(data) > len(clean)
    streams = _pdf_text_streams(data)
    assert b"ignore all previous instructions" in streams


def test_render_file_corrupt_docx_is_not_zip_and_has_no_preview(fx, tmp_path):
    c, r, text = fx
    c.unreadable, c.style.file_format = "corrupt_docx", "docx"
    out = R.render_file(r, text, c, tmp_path)
    f = tmp_path / out.path
    assert f.suffix == ".docx" and not zipfile.is_zipfile(f)
    assert out.preview is None and not (tmp_path / c.candidate_id / R.PREVIEW_NAME).exists()
    assert "corrupt" in out.note


@pytest.mark.parametrize("fmt", ["docx", "txt"])
def test_render_file_docx_txt_get_preview(fx, tmp_path, fmt):
    c, r, text = fx
    c.style.file_format = fmt
    out = R.render_file(r, text, c, tmp_path)
    prev = tmp_path / c.candidate_id / R.PREVIEW_NAME
    assert prev.read_bytes().startswith(b"%PDF") and out.preview == f"{c.candidate_id}/{R.PREVIEW_NAME}"
    if fmt == "txt":
        assert (tmp_path / out.path).read_text().startswith("Testy Fixture")


def test_render_file_rerender_replaces_old_files(fx, tmp_path):
    c, r, text = fx
    c.style.file_format = "docx"
    R.render_file(r, text, c, tmp_path)
    c.style.file_format = "pdf"
    R.render_file(r, text, c, tmp_path)
    names = sorted(p.name for p in (tmp_path / c.candidate_id).iterdir())
    assert len(names) == 1 and names[0].endswith(".pdf")


def test_render_pdf_escapes_markup_in_text(fx):
    c, r, _ = fx
    r["summary"] = "Handled <b>unclosed & weird</i> markup <script>"
    assert R.render_pdf(r, c, "modern").startswith(b"%PDF")


def test_filename_is_deterministic(fx):
    c, _, _ = fx
    assert R.filename_for(c, "pdf") == R.filename_for(c, "pdf")
    assert "/" not in R.filename_for(c, "pdf")
